# justifactu - Automated billing justifications
# Copyright (C) 2026  Aleix Mariné Tena (AleixMT), Carles de la Cuadra, David Romero San Millán (DavidRomeroICIQ)
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <http://www.gnu.org/licenses/>.
import zipfile
from unittest.mock import MagicMock, patch

from pathlib import Path

import openpyxl
import pytest

from justifactu.bank_statement import (
    BankEntry,
    _group_into_rows,
    _parse_amount,
    extract_bank_entries,
    index_statement_entries,
    read_sap_id_ledger,
    Word,
    _merge_commission_rows,
    _strip_autofilters,
    find_year_ledger,
)
from justifactu.SAP_ID import SAP_ID

sap_id = "2034567890"


def _write_ledger_row(sheet, date="", concept="", total="", sap_id="") -> None:
    row = [None] * 10
    row[0] = date
    row[3] = concept
    row[7] = total
    row[9] = sap_id
    sheet.append(row)


def _make_ledger(path: Path, rows: list[dict]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["header"] * 10)  # row 1 - skipped by read_sap_id_ledger's min_row=2
    for row in rows:
        _write_ledger_row(sheet, **row)
    workbook.save(path)


# ── _parse_amount ─────────────────────────────────────────────────────────────


def test_parse_amount_simple():
    assert _parse_amount("56,78") == 56.78


def test_parse_amount_with_thousands_separator():
    assert _parse_amount("1.234,56") == 1234.56


# ── _group_into_rows ──────────────────────────────────────────────────────────


def test_group_into_rows_groups_same_top():

    word_a: Word = {"text": "0101", "x0": 10.0, "top": 100.0}
    word_b: Word = {"text": "Compra", "x0": 50.0, "top": 100.0}
    result = _group_into_rows([word_a, word_b])
    assert result == [[word_a, word_b]]


def test_group_into_rows_separates_different_top():

    word_a: Word = {"text": "0101", "x0": 10.0, "top": 100.0}
    word_b: Word = {"text": "Compra", "x0": 50.0, "top": 150.0}
    result = _group_into_rows([word_a, word_b])
    assert result == [[word_a], [word_b]]


def test_group_into_rows_sorts_row_left_to_right():
    word_a: Word = {"text": "0101", "x0": 10.0, "top": 100.0}
    word_b: Word = {"text": "Compra", "x0": 50.0, "top": 100.0}
    result = _group_into_rows([word_b, word_a])
    assert result == [[word_a, word_b]]


# ── extract_bank_entries ───────────────────────────────────────────────────────


def test_extract_bank_entries_parses_valid_row(tmp_path):
    pdf_path = tmp_path / "statement.pdf"
    pdf_path.touch()

    mock_page = MagicMock()
    mock_page.extract_words.return_value = [
        {"text": "0101", "x0": 0.0, "top": 100.0},
        {"text": "Compra", "x0": 70.0, "top": 100.0},
        {"text": "56,78", "x0": 300.0, "top": 100.0},
    ]
    mock_pdf = MagicMock()
    mock_pdf.pages = [mock_page]
    mock_pdfplumber_open = MagicMock()
    mock_pdfplumber_open.__enter__.return_value = mock_pdf

    with patch(
        "justifactu.bank_statement.pdfplumber.open", return_value=mock_pdfplumber_open
    ):
        result = extract_bank_entries(pdf_path)

    assert result == [BankEntry(date="0101", concept="Compra", total=56.78)]


def test_extract_bank_entries_skips_non_transaction_rows(tmp_path):
    pdf_path = tmp_path / "statement.pdf"
    pdf_path.touch()

    mock_page = MagicMock()
    mock_page.extract_words.return_value = [
        {"text": "01001", "x0": 0.0, "top": 100.0},
        {"text": "Compra", "x0": 70.0, "top": 100.0},
        {"text": "56,78", "x0": 300.0, "top": 100.0},
    ]
    mock_pdf = MagicMock()
    mock_pdf.pages = [mock_page]
    mock_pdfplumber_open = MagicMock()
    mock_pdfplumber_open.__enter__.return_value = mock_pdf

    with patch(
        "justifactu.bank_statement.pdfplumber.open", return_value=mock_pdfplumber_open
    ):
        result = extract_bank_entries(pdf_path)

    assert result == []


def test_extract_bank_entries_skips_row_with_no_amount(tmp_path):
    pdf_path = tmp_path / "statement.pdf"
    pdf_path.touch()

    mock_page = MagicMock()
    mock_page.extract_words.return_value = [
        {"text": "0101", "x0": 0.0, "top": 100.0},
        {"text": "Compra", "x0": 70.0, "top": 100.0},
    ]
    mock_pdf = MagicMock()
    mock_pdf.pages = [mock_page]
    mock_pdfplumber_open = MagicMock()
    mock_pdfplumber_open.__enter__.return_value = mock_pdf

    with patch(
        "justifactu.bank_statement.pdfplumber.open", return_value=mock_pdfplumber_open
    ):
        result = extract_bank_entries(pdf_path)

    assert result == []


def test_extract_bank_entries_excludes_words_past_concept_column(tmp_path):
    pdf_path = tmp_path / "statement.pdf"
    pdf_path.touch()

    mock_page = MagicMock()
    mock_page.extract_words.return_value = [
        {"text": "0101", "x0": 0.0, "top": 100.0},
        {"text": "Compra", "x0": 70.0, "top": 100.0},
        {"text": "Patata", "x0": 290.0, "top": 100.0},
        {"text": "56,78", "x0": 300.0, "top": 100.0},
    ]
    mock_pdf = MagicMock()
    mock_pdf.pages = [mock_page]
    mock_pdfplumber_open = MagicMock()
    mock_pdfplumber_open.__enter__.return_value = mock_pdf

    with patch(
        "justifactu.bank_statement.pdfplumber.open", return_value=mock_pdfplumber_open
    ):
        result = extract_bank_entries(pdf_path)

    assert result == [BankEntry(date="0101", concept="Compra", total=56.78)]


def test_extract_bank_entries_combines_multiple_pages(tmp_path):
    pdf_path = tmp_path / "statement.pdf"
    pdf_path.touch()
    mock_page_1 = MagicMock()
    mock_page_2 = MagicMock()
    mock_page_1.extract_words.return_value = [
        {"text": "0101", "x0": 0.0, "top": 100.0},
        {"text": "Compra", "x0": 70.0, "top": 100.0},
        {"text": "56,78", "x0": 300.0, "top": 100.0},
    ]
    mock_page_2.extract_words.return_value = [
        {"text": "0202", "x0": 0.0, "top": 100.0},
        {"text": "Taxi", "x0": 70.0, "top": 100.0},
        {"text": "82,99", "x0": 300.0, "top": 100.0},
    ]
    mock_pdf = MagicMock()
    mock_pdf.pages = [mock_page_1, mock_page_2]
    mock_pdfplumber_open = MagicMock()
    mock_pdfplumber_open.__enter__.return_value = mock_pdf
    with patch(
        "justifactu.bank_statement.pdfplumber.open", return_value=mock_pdfplumber_open
    ):
        result = extract_bank_entries(pdf_path)
    assert result == [
        BankEntry(date="0101", concept="Compra", total=56.78),
        BankEntry(date="0202", concept="Taxi", total=82.99),
    ]


# ── read_sap_id_ledger ───────────────────────────────────────────────────────


def test_read_sap_id_ledger_maps_date_total_to_sap_id(tmp_path):
    ledger_path = tmp_path / "ledger.xlsx"
    _make_ledger(
        ledger_path,
        [{"date": "01/01", "concept": "Compra", "total": 56.78, "sap_id": sap_id}],
    )

    result = read_sap_id_ledger(ledger_path)

    assert result == {("0101", 56.78): SAP_ID(sap_id)}


def test_read_sap_id_ledger_skips_row_without_sap_id(tmp_path):
    ledger_path = tmp_path / "ledger.xlsx"
    _make_ledger(
        ledger_path,
        [{"date": "01/01", "concept": "Compra", "total": 56.78, "sap_id": None}],
    )

    result = read_sap_id_ledger(ledger_path)

    assert result == {}


def test_read_sap_id_ledger_skips_non_numeric_total(tmp_path):
    ledger_path = tmp_path / "ledger.xlsx"
    _make_ledger(
        ledger_path,
        [{"date": "01/01", "concept": "Compra", "total": "N/A", "sap_id": sap_id}],
    )

    result = read_sap_id_ledger(ledger_path)

    assert result == {}


def test_read_sap_id_ledger_skips_invalid_sap_id(tmp_path):
    ledger_path = tmp_path / "ledger.xlsx"
    _make_ledger(
        ledger_path,
        [
            {
                "date": "01/01",
                "concept": "Compra",
                "total": 56.78,
                "sap_id": "not-a-sap-id",
            }
        ],
    )

    result = read_sap_id_ledger(ledger_path)

    assert result == {}


def test_read_sap_id_ledger_multiple_rows(tmp_path):
    ledger_path = tmp_path / "ledger.xlsx"
    sap_id_a = SAP_ID("2034567890")
    sap_id_b = SAP_ID("2034567891")
    _make_ledger(
        ledger_path,
        [
            {
                "date": "01/01",
                "concept": "Compra",
                "total": 56.78,
                "sap_id": str(sap_id_a),
            },
            {
                "date": "02/02",
                "concept": "Taxi",
                "total": 12.34,
                "sap_id": str(sap_id_b),
            },
        ],
    )

    result = read_sap_id_ledger(ledger_path)

    assert result == {
        ("0101", 56.78): sap_id_a,
        ("0202", 12.34): sap_id_b,
    }


# ── index_statement_entries ───────────────────────────────────────────────────────


def test_index_statement_entries_matches_by_date_and_total(tmp_path):
    pdf_path = tmp_path / "statement.pdf"
    ledger_path = tmp_path / "ledger.xlsx"
    entry = BankEntry(date="0101", concept="Compra", total=56.78)
    sap_id_instance = SAP_ID(sap_id)

    with (
        patch("justifactu.bank_statement.extract_bank_entries", return_value=[entry]),
        patch(
            "justifactu.bank_statement.read_sap_id_ledger",
            return_value={("0101", 56.78): sap_id_instance},
        ),
    ):
        result = index_statement_entries(pdf_path, ledger_path)

    assert result == {sap_id_instance: entry}


def test_index_statement_entries_skips_unmatched_entry(tmp_path):
    pdf_path = tmp_path / "statement.pdf"
    ledger_path = tmp_path / "ledger.xlsx"
    entry = BankEntry(date="0101", concept="Compra", total=56.78)

    with (
        patch("justifactu.bank_statement.extract_bank_entries", return_value=[entry]),
        patch(
            "justifactu.bank_statement.read_sap_id_ledger",
            return_value={("0202", 12.34): SAP_ID(sap_id)},
        ),
    ):
        result = index_statement_entries(pdf_path, ledger_path)

    assert result == {}


def test_index_statement_entries_multiple_entries(tmp_path):
    pdf_path = tmp_path / "statement.pdf"
    ledger_path = tmp_path / "ledger.xlsx"
    entry_a = BankEntry(date="0101", concept="Compra", total=56.78)
    entry_b = BankEntry(date="0202", concept="Taxi", total=12.34)
    sap_id_a = SAP_ID("2034567890")
    sap_id_b = SAP_ID("2034567891")

    with (
        patch(
            "justifactu.bank_statement.extract_bank_entries",
            return_value=[entry_a, entry_b],
        ),
        patch(
            "justifactu.bank_statement.read_sap_id_ledger",
            return_value={
                ("0101", 56.78): sap_id_a,
                ("0202", 12.34): sap_id_b,
            },
        ),
    ):
        result = index_statement_entries(pdf_path, ledger_path)

    assert result == {sap_id_a: entry_a, sap_id_b: entry_b}


# ── _merge_commission_rows ────────────────────────────────────────────────


def test_merge_commission_rows_merges_matching_commission():
    main_entry = BankEntry(date="0105", concept="AAA", total=12.34)
    commission_entry = BankEntry(
        date="0105", concept="Comissió per conversió de divises (*)", total=1.0
    )

    result = _merge_commission_rows([main_entry, commission_entry])

    assert result == [BankEntry(date="0105", concept="AAA", total=13.34)]


def test_merge_commission_rows_leaves_normal_rows_untouched():
    entry_a = BankEntry(date="0101", concept="Compra", total=56.78)
    entry_b = BankEntry(date="0202", concept="Taxi", total=12.34)

    result = _merge_commission_rows([entry_a, entry_b])

    assert result == [entry_a, entry_b]


def test_merge_commission_rows_does_not_merge_across_different_dates():
    entry_a = BankEntry(date="0101", concept="Compra", total=56.78)
    commission_entry = BankEntry(
        date="0202", concept="Comissió per conversió de divises (*)", total=1.0
    )

    result = _merge_commission_rows([entry_a, commission_entry])

    assert result == [entry_a, commission_entry]


# ── find_year_ledger ──────────────────────────────────────────────────────────


def test_find_year_ledger_returns_single_ledger(tmp_path):
    ledger = tmp_path / "Targetes_2026.xlsx"
    ledger.touch()

    assert find_year_ledger(tmp_path) == ledger


def test_find_year_ledger_none_when_missing(tmp_path):
    assert find_year_ledger(tmp_path) is None


def test_find_year_ledger_none_when_more_than_one(tmp_path):
    (tmp_path / "a.xlsx").touch()
    (tmp_path / "b.xlsx").touch()

    assert find_year_ledger(tmp_path) is None


def test_find_year_ledger_ignores_excel_lock_file(tmp_path):
    ledger = tmp_path / "Targetes_2026.xlsx"
    ledger.touch()
    (tmp_path / "~$Targetes_2026.xlsx").touch()

    assert find_year_ledger(tmp_path) == ledger


# ── _strip_autofilters ────────────────────────────────────────────────────────


def _add_whole_row_autofilter(path: Path) -> None:
    """Rewrites the sheet's AutoFilter to a whole-row range (ref="1:644"),
    the shape Excel writes and openpyxl refuses to parse."""
    workbook = openpyxl.load_workbook(path)
    workbook.active.auto_filter.ref = "A1:J2"
    workbook.save(path)

    rewritten = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(rewritten, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("xl/worksheets/"):
                data = data.replace(
                    b'<autoFilter ref="A1:J2"', b'<autoFilter ref="1:644"'
                )
            dst.writestr(item, data)
    rewritten.replace(path)


def test_whole_row_autofilter_breaks_plain_openpyxl(tmp_path):
    ledger_path = tmp_path / "ledger.xlsx"
    _make_ledger(
        ledger_path,
        [{"date": "01/01", "concept": "Compra", "total": 56.78, "sap_id": sap_id}],
    )
    _add_whole_row_autofilter(ledger_path)

    with pytest.raises(ValueError):
        openpyxl.load_workbook(ledger_path, data_only=True)


def test_read_sap_id_ledger_reads_ledger_with_whole_row_autofilter(tmp_path):
    ledger_path = tmp_path / "ledger.xlsx"
    _make_ledger(
        ledger_path,
        [{"date": "01/01", "concept": "Compra", "total": 56.78, "sap_id": sap_id}],
    )
    _add_whole_row_autofilter(ledger_path)

    assert read_sap_id_ledger(ledger_path) == {("0101", 56.78): SAP_ID(sap_id)}


def test_strip_autofilters_leaves_original_file_untouched(tmp_path):
    ledger_path = tmp_path / "ledger.xlsx"
    _make_ledger(
        ledger_path,
        [{"date": "01/01", "concept": "Compra", "total": 56.78, "sap_id": sap_id}],
    )
    _add_whole_row_autofilter(ledger_path)
    before = ledger_path.read_bytes()

    _strip_autofilters(ledger_path)

    assert ledger_path.read_bytes() == before
